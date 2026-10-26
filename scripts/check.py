"""Check your homework before you send it.

  python scripts/check.py list        # what can be checked
  python scripts/check.py hw1.2       # check one homework

It runs your files against the Nur Finance data and says, task by task, whether your result is right.
It never shows the right answer: checks/checks.json holds only fingerprints of the answers.
A green check means "your number matches"; your mentor still reads the queries, the comments and the text.

SQL: write SQLite or SQL Server; the common SQL Server words (TOP, ISNULL, LEN, GETDATE, [brackets], dbo.) are
translated. If a query still cannot run here, the checker reads the number in your "-- result:" line instead.
Needs: Python 3.10+, and for some homework openpyxl / pandas / jupyter (pip install -r requirements.txt).
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, "data", "raw", "nur.db")
SPEC = json.load(open(os.path.join(ROOT, "checks", "checks.json"), encoding="utf-8"))
SALT = SPEC["salt"]
GREEN, RED, GREY = "OK   ", "NO   ", "--   "


# ---------------------------------------------------------------- matching numbers without knowing them
def fp(x: float, dp: int) -> str:
    return hashlib.sha256((SALT + f"{round(x, dp):.{dp}f}").encode()).hexdigest()[:24]


def to_num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", "").rstrip("%")
    try:
        return float(s)
    except ValueError:
        return None


class Bag:
    """All numbers seen in a result; answers whether an expected fingerprint is among them."""

    def __init__(self, values=()):
        self.nums = [n for n in (to_num(v) for v in values) if n is not None]

    def add_text(self, text: str):
        for tok in re.findall(r"-?\d[\d  ,]*\.?\d*%?", text):
            n = to_num(tok)
            if n is not None:
                self.nums.append(n)

    def has(self, exp: dict) -> bool:
        dp = exp["dp"]
        for n in self.nums:
            if fp(n, dp) == exp["h"] or fp(n * 100, dp) == exp["h"]:  # 0.603 and 60.3 both count
                return True
        return False


# ---------------------------------------------------------------- running SQL
def translate(sql: str) -> str:
    s = re.sub(r"\[([^\]]+)\]", r'"\1"', sql)
    s = re.sub(r'"?\bdbo"?\.', "", s, flags=re.I)
    s = re.sub(r"\bYEAR\s*\(([^()]+)\)", r"CAST(strftime('%Y', \1) AS INTEGER)", s, flags=re.I)
    s = re.sub(r"\bMONTH\s*\(([^()]+)\)", r"CAST(strftime('%m', \1) AS INTEGER)", s, flags=re.I)
    s = re.sub(r"\bDAY\s*\(([^()]+)\)", r"CAST(strftime('%d', \1) AS INTEGER)", s, flags=re.I)
    s = re.sub(r"\bISNULL\s*\(", "IFNULL(", s, flags=re.I)
    s = re.sub(r"\bLEN\s*\(", "LENGTH(", s, flags=re.I)
    s = re.sub(r"\bGETDATE\s*\(\s*\)", "'2026-10-01'", s, flags=re.I)
    m = re.search(r"\bSELECT\s+(DISTINCT\s+)?TOP\s*\(?\s*(\d+)\s*\)?", s, flags=re.I)
    if m:
        s = s[:m.start()] + "SELECT " + (m.group(1) or "") + s[m.end():]
        s = s.rstrip().rstrip(";") + f" LIMIT {m.group(2)}"
    return s


def strip_comments(sql: str) -> str:
    sql = re.sub(r"/\*.*?\*/", " ", sql, flags=re.S)
    return "\n".join(line.split("--")[0] for line in sql.splitlines())


def run_block(con, text: str):
    """Run every statement in a block. Returns (cells, rows_of_last_select, error_or_None)."""
    cells, last_rows, err = [], None, None
    for stmt in [s for s in strip_comments(text).split(";") if s.strip()]:
        if re.match(r"\s*(USE|GO)\b", stmt, flags=re.I):
            continue
        if re.match(r"\s*(INSERT|UPDATE|DELETE|DROP|ALTER|TRUNCATE)\b", stmt, flags=re.I):
            continue  # never change the course data
        for attempt in (stmt, translate(stmt)):
            try:
                rows = con.execute(attempt).fetchall()
                cells += [v for r in rows for v in r]
                last_rows = len(rows)
                err = None
                break
            except sqlite3.Error as e:
                err = str(e)
    return cells, last_rows, err


def result_lines(text: str) -> str:
    return " ".join(m for m in re.findall(r"--\s*result\s*:(.*)", text, flags=re.I))


def connect():
    con = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
    return con


def check_sql(c: dict) -> int:
    path = os.path.join(ROOT, c["file"])
    if not os.path.exists(path):
        print(f"{RED}file not found: {c['file']}")
        return 1
    text = open(path, encoding="utf-8", errors="ignore").read()
    parts = re.split(r"--\s*Q(\d+)\s*[:.]", text, flags=re.I)
    # each block starts with the rest of the "-- Qn: question" line, which is comment text: drop it
    blocks = {f"Q{parts[i]}": parts[i + 1].split("\n", 1)[1] if "\n" in parts[i + 1] else ""
              for i in range(1, len(parts) - 1, 2)}
    con, bad = connect(), 0
    for task, exp in c["tasks"].items():
        block = blocks.get(task)
        if block is None:
            print(f"{RED}{task}: no '-- {task}:' line found in the file")
            bad += 1
            continue
        cells, nrows, err = run_block(con, block)
        bag = Bag(cells)
        ok_vals = all(bag.has(e) for e in exp.get("values", []))
        ok_rows = "rows" not in exp or (nrows is not None and Bag([nrows]).has(exp["rows"]))
        if ok_vals and ok_rows and (cells or nrows is not None):
            print(f"{GREEN}{task}")
            continue
        fb = Bag()
        fb.add_text(result_lines(block))
        if err and all(fb.has(e) for e in exp.get("values", [])) and ("rows" not in exp or fb.has(exp["rows"])):
            print(f"{GREEN}{task} (from your '-- result:' line; the query did not run here: {err})")
            continue
        bad += 1
        why = []
        if err:
            why.append(f"the query did not run here ({err})")
        if not ok_vals:
            why.append("a number in the result is not the expected one")
        if not ok_rows:
            why.append(f"the number of rows ({nrows}) is not the expected one")
        print(f"{RED}{task}: " + "; ".join(why))
    return bad


def check_sql_files(c: dict) -> int:
    con, bad = connect(), 0
    for rel, exp in c["files"].items():
        path = os.path.join(ROOT, rel)
        name = os.path.basename(rel)
        if not os.path.exists(path):
            print(f"{RED}{name}: file not found ({rel})")
            bad += 1
            continue
        text = open(path, encoding="utf-8", errors="ignore").read()
        cells, _, err = run_block(con, text)
        bag = Bag(cells)
        bag.add_text(result_lines(text))
        found = sum(bag.has(e) for e in exp["values"])
        if found == len(exp["values"]):
            print(f"{GREEN}{name}: all {found} key numbers found")
        else:
            bad += 1
            print(f"{RED}{name}: {found} of {len(exp['values'])} key numbers found" + (f" (last error: {err})" if err else ""))
    for name, what in c.get("notes", {}).items():
        if not name.startswith("_"):
            print(f"{GREY}{name}: {what}")
    return bad


# ---------------------------------------------------------------- Excel
def check_xlsx(c: dict) -> int:
    try:
        import openpyxl
    except ImportError:
        print("Install openpyxl first: pip install openpyxl")
        return 1
    path = os.path.join(ROOT, c["file"])
    if not os.path.exists(path):
        print(f"{RED}file not found: {c['file']}")
        return 1
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    cells = [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    if not any(isinstance(v, (int, float)) for v in cells):
        print(f"{RED}no calculated values in the file: open it in Excel, let it calculate, save, and run again")
        return 1
    bag, bad = Bag(cells), 0
    for label, exp in c["values"].items():
        if bag.has(exp):
            print(f"{GREEN}{label}")
        else:
            bad += 1
            print(f"{RED}{label}: not found among your cells")
    return bad


def check_xlsx_tables(c: dict) -> int:
    try:
        import openpyxl
    except ImportError:
        print("Install openpyxl first: pip install openpyxl")
        return 1
    path = os.path.join(ROOT, c["file"])
    if not os.path.exists(path):
        print(f"{RED}file not found: {c['file']}")
        return 1
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    bad = 0
    for name, exp in c["tables"].items():
        ws = next((w for w in wb.worksheets if name.lower() in w.title.lower().replace(" ", "_")), None)
        if ws is None:
            bad += 1
            print(f"{RED}{name}: no sheet with this name")
            continue
        rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
        n = len(rows) - 1
        msg, ok = [], Bag([n]).has(exp["rows"])
        msg.append(f"{n} data rows " + ("(right)" if ok else "(not the expected count)"))
        if "sum" in exp:
            head = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
            if "amount" in head:
                k = head.index("amount")
                s = sum(to_num(r[k]) or 0 for r in rows[1:])
                sok = Bag([s]).has(exp["sum"])
                ok = ok and sok
                msg.append("amount total " + ("right" if sok else "not the expected total"))
            else:
                ok = False
                msg.append("no 'amount' column")
        bad += 0 if ok else 1
        print(f"{GREEN if ok else RED}{name}: " + ", ".join(msg))
    return bad


# ---------------------------------------------------------------- Python
def notebook_home(file: str) -> str:
    """The data/python folder the starter notebook lives in, so that its relative paths (DATA = Path("..") or
    Path("../..")) resolve exactly as they did when the learner ran it there."""
    name = os.path.basename(file)
    m = re.search(r"week-0*(\d+)", file)
    folders = [os.path.join(ROOT, "data", "python", f"week{m.group(1)}")] if m else []
    folders.append(os.path.join(ROOT, "data", "python"))
    for folder in folders:
        if os.path.exists(os.path.join(folder, name)):
            return folder
    return os.path.dirname(os.path.join(ROOT, file))


def check_notebook(c: dict) -> int:
    path = os.path.join(ROOT, c["file"])
    if not os.path.exists(path):
        print(f"{RED}file not found: {c['file']}")
        return 1
    out = os.path.join(tempfile.mkdtemp(), "run.ipynb")
    home = notebook_home(c["file"])
    # run a temporary copy next to the starter notebook, so its "../.." still means data/; removed afterwards
    run_copy = os.path.join(home, f"_check_{os.getpid()}_{os.path.basename(path)}")
    print("running your notebook from the top (this can take a minute)…")
    try:
        with open(path, "rb") as src, open(run_copy, "wb") as dst:
            dst.write(src.read())
        r = subprocess.run([sys.executable, "-m", "jupyter", "nbconvert", "--to", "notebook", "--execute",
                            "--allow-errors", "--ExecutePreprocessor.timeout=600", "--output", out, run_copy],
                           cwd=home, capture_output=True, text=True)
    finally:
        if os.path.exists(run_copy):
            os.remove(run_copy)
    if r.returncode != 0 or not os.path.exists(out):
        print(f"{RED}the notebook could not be run: {r.stderr.strip().splitlines()[-1] if r.stderr.strip() else 'unknown error'}")
        return 1
    nb = json.load(open(out, encoding="utf-8"))
    errors = []
    for i, cell in enumerate(nb["cells"], 1):
        for o in cell.get("outputs", []):
            if o.get("output_type") == "error":
                errors.append((i, o.get("ename"), o.get("evalue", "")[:120]))
    if not errors:
        print(f"{GREEN}every cell ran and every assert passed")
        return 0
    for i, ename, ev in errors:
        print(f"{RED}cell {i}: {ename} {ev}")
    print(f"{len(errors)} cell(s) failed. A NotImplementedError means a TODO is still open.")
    return len(errors)


def check_csv(c: dict) -> int:
    script = os.path.join(ROOT, c["run"])
    folder = os.path.dirname(script)
    if os.path.exists(script):
        print(f"running {c['run']}…")
        r = subprocess.run([sys.executable, script], cwd=folder, capture_output=True, text=True)
        if r.returncode != 0:
            tail = (r.stderr.strip().splitlines() or ["unknown error"])[-1]
            print(f"{RED}the script stopped with an error: {tail}")
            return 1
    else:
        print(f"{GREY}{c['run']} not found; checking the CSV files already in {os.path.relpath(folder, ROOT)}")
    bad = 0
    for name, exp in c["files"].items():
        path = os.path.join(folder, name)
        if not os.path.exists(path):
            bad += 1
            print(f"{RED}{name}: not written")
            continue
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        ok, msg = Bag([len(rows)]).has(exp["rows"]), [f"{len(rows)} rows"]
        if "sum" in exp:
            s = sum(to_num(r.get(exp["sum"]["col"])) or 0 for r in rows)
            sok = Bag([s]).has(exp["sum"])
            ok = ok and sok
            msg.append(f"{exp['sum']['col']} total " + ("right" if sok else "not the expected total"))
        if "positive" in exp:
            p = sum(1 for r in rows if (to_num(r.get(exp["positive"]["col"])) or 0) > 0)
            pok = Bag([p]).has(exp["positive"])
            ok = ok and pok
            msg.append(f"{p} rows with {exp['positive']['col']} > 0 " + ("(right)" if pok else "(not the expected count)"))
        bad += 0 if ok else 1
        print(f"{GREEN if ok else RED}{name}: " + ", ".join(msg))
    return bad


KINDS = {"sql": check_sql, "sql-files": check_sql_files, "xlsx": check_xlsx, "xlsx-tables": check_xlsx_tables,
         "notebook": check_notebook, "csv": check_csv}


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    arg = sys.argv[1].lower() if len(sys.argv) > 1 else "list"
    checks = SPEC["checks"]
    if arg == "list" or arg not in checks:
        if arg != "list":
            print(f"'{arg}' is not a homework the checker knows.\n")
        print("Homework you can check:")
        for k, c in checks.items():
            print(f"  python scripts/check.py {k:<6}  {c['title']}")
        print("\nPower BI, memos, CVs and the capstone are reviewed by your mentor.")
        return
    c = checks[arg]
    print(f"\n{c['title']}\n" + "-" * len(c["title"]))
    bad = KINDS[c["kind"]](c)
    print("\nAll checks passed — send it." if bad == 0 else f"\n{bad} check(s) to fix. Fix, run again, then send.")
    sys.exit(0 if bad == 0 else 1)


if __name__ == "__main__":
    main()
